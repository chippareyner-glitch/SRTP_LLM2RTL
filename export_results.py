"""Export real DeepSeek experiment evaluations using the project's openpyxl workflow."""
import argparse
import csv
import json
from datetime import datetime
from pathlib import Path
import shutil
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

HEADERS = ['题号', '类别', '迭代次数', '语法是否通过', '功能是否通过', '报错原因']
ROOT = Path(__file__).resolve().parent

def export_report(run_dir, output=None, categories=None):
    run_dir = Path(run_dir).resolve()
    output = Path(output or run_dir/'数据.xlsx').resolve()
    data = json.loads((run_dir/'summary.json').read_text(encoding='utf-8'))
    category_path = Path(categories or ROOT/'case_categories.json')
    mapping = json.loads(category_path.read_text(encoding='utf-8')) if category_path.exists() else {}
    rows = []
    for result in data['results']:
        cid = result['case_id']
        attempts = result.get('attempts', [])
        # Count started generation rounds, including a failed API request.
        rounds = max(len(attempts), len(list((run_dir/cid).glob('attempt_*/prompt.txt'))))
        evaluation = attempts[-1] if attempts else {}
        if result['status'] == 'error':
            syntax = function = '未判定'
            reason = 'API/生成异常：' + result.get('error', '')
        else:
            syntax = {True:'通过', False:'未通过'}.get(evaluation.get('syntax_pass'), '未判定')
            function = '通过' if evaluation.get('func_pass') else ('未执行' if syntax != '通过' else '未通过')
            reason = '' if evaluation.get('func_pass') else evaluation.get('error_text') or result['status']
        rows.append([cid, mapping.get(cid, '未分类'), rounds, syntax, function, reason[:32000]])
    if len({r[0] for r in rows}) != len(rows):
        raise ValueError('Duplicate case IDs in experiment summary')
    wb = Workbook(); ws = wb.active; ws.title = 'Test Results'
    ws.append(HEADERS)
    for row in rows: ws.append(row)
    ws.freeze_panes = 'A2'
    for cell in ws[1]:
        cell.font = Font(name='微软雅黑', bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='253B60')
        cell.alignment = Alignment(vertical='center')
    ws.row_dimensions[1].height = 28
    for col, width in zip('ABCDEF', [36,30,14,18,18,110]): ws.column_dimensions[col].width = width
    for cells in ws.iter_rows(min_row=2):
        for cell in cells:
            cell.font = Font(name='微软雅黑', size=11)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        ws.row_dimensions[cells[0].row].height = 32 if not cells[5].value else min(350, max(65, len(str(cells[5].value).splitlines())*16))
    if rows:
        table = Table(displayName='TestResults', ref=f'A1:F{len(rows)+1}')
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        ws.add_table(table)
    output.parent.mkdir(parents=True, exist_ok=True)
    if output.exists():
        backup = output.with_name(output.stem+'_备份_'+datetime.now().strftime('%Y%m%d_%H%M%S_%f')+output.suffix)
        shutil.copy2(output, backup)
    temporary = output.with_name(output.stem+'.tmp.xlsx')
    wb.save(temporary)
    check = load_workbook(temporary, read_only=True, data_only=True)
    assert list(check.active.values) == [tuple(HEADERS)] + [tuple(v if v != '' else None for v in r) for r in rows]
    check.close()
    try:
        temporary.replace(output)
    except PermissionError:
        output = output.with_name(output.stem+'_六列_'+datetime.now().strftime('%Y%m%d_%H%M%S')+output.suffix)
        temporary.replace(output)
        print('Requested workbook is in use; saved to a new file instead.')
    with (run_dir/'results.csv').open('w', encoding='utf-8-sig', newline='') as f:
        writer=csv.writer(f); writer.writerow(HEADERS); writer.writerows(rows)
    print(f'Excel: {output} ({len(rows)} cases)')
    return output

if __name__ == '__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--run', type=Path, required=True)
    parser.add_argument('--output', type=Path)
    parser.add_argument('--categories', type=Path)
    args=parser.parse_args()
    export_report(args.run, args.output, args.categories)
